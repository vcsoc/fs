from pathlib import Path
import time

from docx import Document
from pypdf import PdfWriter

from fs.core import Store, extract
from fs.metadata import birth_time


def test_docx_and_pdf(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    document = Document()
    document.add_paragraph("Searchable proposal for a solar project")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Budget estimate"
    document.save(root / "proposal.docx")
    assert "solar project" in extract(root / "proposal.docx")
    assert "Budget estimate" in extract(root / "proposal.docx")
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with (root / "blank.pdf").open("wb") as stream:
        writer.write(stream)
    store = Store(tmp_path / "db")
    store.save_preferences({"scan_duty_percent": 100})
    store.index(root, semantic=False)
    assert store.search("solar", mode="text")["results"][0]["name"] == "proposal.docx"
    assert store.status()["errors"] == 0


def test_creation_time_does_not_use_unix_ctime(tmp_path):
    path = Path(tmp_path) / "new.txt"
    path.write_text("new document")
    stat = path.stat()
    created = birth_time(path, stat)
    if created is not None:
        assert abs(created - time.time()) < 10
    # Unsupported filesystems legitimately expose no creation date.
    assert created is None or isinstance(created, float)
