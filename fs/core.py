from __future__ import annotations

import codecs
import json
import os
import sqlite3
import sys
import threading
import tempfile
import time
import urllib.error
import urllib.request
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

import numpy as np
import regex
from .locking import writer_lock
from .metadata import birth_time
from .preferences import validate, worker_layout
from .sources import SourceMonitor, virtual_source, volume_identity, ensure_gvfs_bridge

DEFAULT_MODEL = "embeddinggemma-2:740m"
EXCLUDED = {".git", ".venv", "venv", "node_modules", "__pycache__", ".cache", "$RECYCLE.BIN", "System Volume Information"}
TEXT_EXTENSIONS = set("txt md rst csv tsv log json jsonl xml html htm css scss js jsx ts tsx py rs go java c h cpp hpp sh bash zsh fish toml yaml yml ini cfg conf sql tex vue svelte rb php swift kt dart r env gitignore dockerfile makefile".split())
MAX_BYTES = 8 * 1024 * 1024
MAX_CHUNKS = 64
VIRTUAL_ROOTS = tuple(Path(p) for p in ("/proc", "/sys", "/dev", "/run")) if sys.platform.startswith("linux") else ()


def data_directory() -> Path:
    if os.name == "nt":
        return Path(os.environ.get("LOCALAPPDATA", Path.home())) / "fs"
    if __import__("sys").platform == "darwin":
        return Path.home() / "Library/Application Support/fs"
    return Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local/share")) / "fs"


class Embedder:
    def __init__(self, model=None, url="http://127.0.0.1:11434"):
        self.model_override = model is not None
        self.model = model or DEFAULT_MODEL
        self.url = url.rstrip("/")
        self.options = {"num_thread": 2, "num_gpu": 0}
        self.directory = None
        self.backend = "auto"
        self.native = None

    @property
    def provider(self):
        from .native_model import DIRECT_MODELS, FULL_MODELS
        if self.directory and self.model in DIRECT_MODELS:
            if self.backend == "direct":
                return "direct"
            try:
                state = json.loads((Path(self.directory) / "runtime.json").read_text())
                if state.get("model") == self.model and state.get("provider") in {"direct", "ollama"}:
                    return state["provider"]
            except (OSError, ValueError):
                pass
            from .native_model import NativeModel, FULL_MODELS
            if self.backend == "auto" and NativeModel(self.directory, full=self.model in FULL_MODELS).ready():
                return "direct"
        return "ollama"

    @property
    def identity(self):
        if self.provider == "direct":
            from .native_model import IDENTITY
            return IDENTITY
        return self.model + ":ollama"

    @property
    def supports_media(self):
        from .native_model import FULL_MODELS
        return self.model in FULL_MODELS and self.provider == 'direct'

    def embed_media(self, **inputs):
        from .native_model import NativeModel
        threads = self.options.get('num_thread', 2)
        if self.native is None or self.native.threads != threads or not self.native.full:
            self.native = NativeModel(self.directory, threads, full=True)
        return self.native.embed_media(**inputs)

    def embed(self, texts: list[str], query=False) -> np.ndarray:
        inputs = [f"task: search result | query: {text}" if query else text for text in texts]
        if self.provider == "direct":
            from .native_model import NativeModel, FULL_MODELS
            threads = self.options.get("num_thread", 2)
            if self.native is None or self.native.threads != threads or self.native.full != (self.model in FULL_MODELS):
                self.native = NativeModel(self.directory, threads, full=self.model in FULL_MODELS)
            return self.native.embed(inputs)
        request = urllib.request.Request(
            self.url + "/api/embed",
            data=json.dumps({"model": self.model, "input": inputs, "truncate": True, "keep_alive": "10m", "options": self.options}).encode(),
            headers={"Content-Type": "application/json"},
        )
        try:
            with urllib.request.urlopen(request, timeout=120) as response:
                values = np.asarray(json.load(response)["embeddings"], dtype=np.float32)
        except (OSError, ValueError, KeyError) as exc:
            detail = str(exc)
            if isinstance(exc, urllib.error.HTTPError):
                detail = exc.read(4096).decode(errors="replace")
                if "MLX" in detail:
                    detail += " This model variant requires MLX, which this Ollama runtime cannot load. Choose a compatible model explicitly in Preferences or install a supported runtime."
            raise RuntimeError(f"Embedding service unavailable for {self.model}. Start Ollama and pull the selected model. {detail[:1500]}") from exc
        if values.ndim != 2 or len(values) != len(texts) or not np.isfinite(values).all():
            raise RuntimeError("The model returned invalid embeddings.")
        norms = np.linalg.norm(values, axis=1, keepdims=True)
        if np.any(norms == 0):
            raise RuntimeError("The model returned an empty embedding.")
        return values / norms


def extract(path: Path, max_bytes=MAX_BYTES) -> str:
    from .media import captions, kind
    if kind(path) in {'video', 'audio'}:
        return '\n'.join(cue['text'] for cue in captions(path))[:100_000]
    if max_bytes and path.stat().st_size > max_bytes:
        return ""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages[:100])[:100_000]
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])[:100_000]
    if suffix.lstrip(".") not in TEXT_EXTENSIONS and path.name.lower() not in TEXT_EXTENSIONS and suffix:
        return ""
    with path.open("rb") as stream:
        raw = stream.read(400_000)
    if b"\x00" in raw[:4096]:
        return ""
    try:
        return codecs.getincrementaldecoder("utf-8-sig")().decode(raw, final=False)[:100_000]
    except UnicodeDecodeError:
        return ""


def chunks(text: str, maximum=MAX_CHUNKS) -> list[str]:
    return [text[i:i + 1400] for i in range(0, len(text), 1200)][:maximum]


def date_bound(value: str | None, upper=False) -> float | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if upper and len(value) == 10:
        parsed += timedelta(days=1)
    return parsed.timestamp()


class Store:
    def __init__(self, directory: Path | str | None = None, embedder: Embedder | None = None):
        self.directory = Path(directory) if directory else data_directory()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.path = self.directory / "index.sqlite3"
        self.embedder = embedder or Embedder()
        self.index_lock = threading.Lock()
        self.cache_lock = threading.Lock()
        self.vector_cache = None
        self.sources = SourceMonitor()
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                INSERT OR IGNORE INTO meta VALUES('revision', '0');
                CREATE TABLE IF NOT EXISTS files(
                    id INTEGER PRIMARY KEY, path TEXT UNIQUE NOT NULL, root TEXT NOT NULL,
                    name TEXT NOT NULL, size INTEGER NOT NULL, modified REAL NOT NULL,
                    created REAL, fingerprint TEXT NOT NULL, content TEXT NOT NULL, error TEXT
                );
                CREATE TABLE IF NOT EXISTS locations(root TEXT PRIMARY KEY);
                INSERT OR IGNORE INTO locations(root) SELECT DISTINCT root FROM files;
                CREATE INDEX IF NOT EXISTS files_root ON files(root);
                CREATE INDEX IF NOT EXISTS files_modified ON files(modified);
                CREATE INDEX IF NOT EXISTS files_created ON files(created);
                CREATE VIRTUAL TABLE IF NOT EXISTS search USING fts5(name, path, content, content='files', content_rowid='id');
                CREATE TRIGGER IF NOT EXISTS files_ai AFTER INSERT ON files BEGIN
                    INSERT INTO search(rowid,name,path,content) VALUES(new.id,new.name,new.path,new.content);
                END;
                CREATE TRIGGER IF NOT EXISTS files_ad AFTER DELETE ON files BEGIN
                    INSERT INTO search(search,rowid,name,path,content) VALUES('delete',old.id,old.name,old.path,old.content);
                END;
                CREATE TRIGGER IF NOT EXISTS files_au AFTER UPDATE ON files BEGIN
                    INSERT INTO search(search,rowid,name,path,content) VALUES('delete',old.id,old.name,old.path,old.content);
                    INSERT INTO search(rowid,name,path,content) VALUES(new.id,new.name,new.path,new.content);
                END;
                CREATE TABLE IF NOT EXISTS favourites(path TEXT PRIMARY KEY, kind TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS media_segments(file_id INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
                    part INTEGER NOT NULL, start REAL NOT NULL, end REAL NOT NULL, text TEXT NOT NULL, PRIMARY KEY(file_id,part));
                CREATE TABLE IF NOT EXISTS vectors(file_id INTEGER NOT NULL REFERENCES files(id) ON DELETE CASCADE,
                    part INTEGER NOT NULL, text TEXT NOT NULL, embedding BLOB NOT NULL, PRIMARY KEY(file_id,part));
            """)

            file_columns = {row[1] for row in db.execute('PRAGMA table_info(files)')}
            for column in ('indexed_at', 'analyzed_at'):
                if column not in file_columns:
                    db.execute('ALTER TABLE files ADD COLUMN ' + column + ' REAL')
            location_columns = {row[1] for row in db.execute('PRAGMA table_info(locations)')}
            if 'device' not in location_columns:
                db.execute('ALTER TABLE locations ADD COLUMN device INTEGER')
            if 'volume' not in location_columns:
                db.execute('ALTER TABLE locations ADD COLUMN volume TEXT')
            if 'total_files' not in location_columns:
                db.execute('ALTER TABLE locations ADD COLUMN total_files INTEGER')
        self.apply_preferences()

    def add_location(self, root):
        path = Path(root).expanduser()
        if not path.is_absolute():
            raise ValueError('Use an absolute folder or mounted-source path.')
        if VIRTUAL_ROOTS and virtual_source(path):
            raise ValueError('Virtual system folders are not indexable sources.')
        resolved = str(path.resolve())
        with self.connect() as db:
            db.execute('INSERT OR IGNORE INTO locations(root) VALUES(?)', (resolved,))
        self.sources.invalidate(resolved)
        return resolved

    def preferences(self):
        with self.connect() as db:
            raw = self._meta(db, "preferences")
            search = self._meta(db, 'search_preferences')
        return validate({**(json.loads(raw) if raw else {}), **(json.loads(search) if search else {})})

    def apply_preferences(self):
        prefs = self.preferences()
        if not getattr(self.embedder, "model_override", True):
            self.embedder.model = prefs["model"]
        self.embedder.directory = self.directory
        self.embedder.backend = prefs["embedding_backend"]
        self.embedder.options = {"num_thread": prefs["cpu_threads"]}
        if not prefs["use_gpu"]:
            self.embedder.options["num_gpu"] = 0
        return prefs

    def save_preferences(self, values):
        prefs = validate({**self.preferences(), **values})
        with self.connect() as db:
            # Keep new search/display settings separate so older scanning
            # workers can still read their resource preferences during upgrades.
            search_keys = {'accuracy', 'result_layout'}
            db.execute("INSERT OR REPLACE INTO meta VALUES('preferences',?)", (json.dumps({key: value for key, value in prefs.items() if key not in search_keys}),))
            db.execute("INSERT OR REPLACE INTO meta VALUES('search_preferences',?)", (json.dumps({key: prefs[key] for key in search_keys}),))
        self.apply_preferences()
        return prefs

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=30)
        try:
            db.row_factory = sqlite3.Row
            db.execute("PRAGMA journal_mode=WAL")
            db.execute("PRAGMA foreign_keys=ON")
            with db:
                yield db
        finally:
            db.close()

    def status(self):
        prefs = self.apply_preferences()
        with self.connect() as db:
            count = db.execute("SELECT count(*),sum(length(content)>0),sum(error IS NOT NULL) FROM files").fetchone()
            return {"files": count[0], "readable": count[1] or 0, "errors": count[2] or 0,
                    "chunks": db.execute("SELECT count(*) FROM vectors").fetchone()[0],
                    "roots": [{**dict(r), 'availability': self.sources.status(r['root'], r['device'], r['volume'])} for r in db.execute("SELECT l.root,l.device,l.volume,l.total_files,count(f.id) AS files FROM locations l LEFT JOIN files f ON f.root=l.root GROUP BY l.root ORDER BY l.root")],
                    "model": self.embedder.model, "database": str(self.path),
                    "indexed_model": self._meta(db, "model"), "preferences": prefs}

    @staticmethod
    def _meta(db, key):
        row = db.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
        return row[0] if row else None

    @staticmethod
    def _bump(db):
        db.execute("UPDATE meta SET value=CAST(value AS INTEGER)+1 WHERE key='revision'")

    def index(self, root: str | Path, semantic=True, progress: Callable | None = None, cancel: threading.Event | None = None):
        if not self.index_lock.acquire(blocking=False):
            raise RuntimeError("An indexing job is already running.")
        try:
            with writer_lock(self.directory / "writer.lock"):
                return self._index(root, semantic, progress or (lambda _: None), cancel or threading.Event())
        finally:
            self.index_lock.release()

    def _index(self, root, semantic, progress, cancel):
        prefs = self.apply_preferences()
        workers, inference_threads = worker_layout(prefs, getattr(self.embedder, 'supports_media', False))
        prefs = {**prefs, 'cpu_threads': inference_threads}
        self.embedder.options['num_thread'] = inference_threads
        root = Path(root).expanduser().resolve()
        ensure_gvfs_bridge(root)
        if not root.is_dir() or (VIRTUAL_ROOTS and virtual_source(root)):
            raise ValueError("Select an existing directory or mounted drive, not a virtual system folder.")
        # Validate the service before changing the index; never mix model spaces.
        if semantic:
            self.embedder.embed(["title: fs | text: file search"])
        with self.connect() as db:
            db.execute("INSERT OR IGNORE INTO locations(root) VALUES(?)", (str(root),))
            previous = db.execute('SELECT device,volume FROM locations WHERE root=?', (str(root),)).fetchone()
            current_device, current_volume = root.stat().st_dev, volume_identity(root)
            if (previous['volume'] is not None and previous['volume'] != current_volume) or (previous['volume'] is None and previous['device'] is not None and previous['device'] != current_device):
                raise ValueError('Source volume identity changed or is unmounted. Cached results are preserved; remove and re-add only if this is intentionally a different source.')
            db.execute('UPDATE locations SET device=?,volume=? WHERE root=?', (current_device, current_volume, str(root)))
            model = self._meta(db, "model")
            if semantic and model != getattr(self.embedder, "identity", self.embedder.model):
                db.execute("DELETE FROM vectors")
                db.execute("INSERT OR REPLACE INTO meta VALUES('model',?)", (getattr(self.embedder, "identity", self.embedder.model),))
                self._bump(db)
            existing = {r["path"]: dict(r) for r in db.execute("SELECT id,path,fingerprint,error FROM files WHERE root=?", (str(root),))}
            embedded = {r[0] for r in db.execute("SELECT DISTINCT file_id FROM vectors")}
        state = {"scanned": 0, "processed": 0, "file_workers": workers, "inference_threads": inference_threads, "updated": 0, "unchanged": 0, "errors": 0, "removed": 0, "path": "", "cancelled": False}
        seen = set()
        walk_errors = []
        def report_progress(payload):
            if payload.get('phase') == 'indexing' and 'total' in payload:
                with self.connect() as db:
                    db.execute('UPDATE locations SET total_files=? WHERE root=?', (payload['total'], str(Path(root).expanduser().resolve())))
            state.update({key: value for key, value in payload.items() if key != 'state'})
            progress(dict(state))
        from .analysis_pool import analyse_file, run_files
        from .media import caption_fingerprint
        def submitted(path):
            seen.add(str(path))
            state['scanned'] += 1
            state['path'] = str(path)
            progress(dict(state))
        def report_path(path):
            state['path'] = path
            progress(dict(state))
        paths = self._paths(root, walk_errors, cancel, report_progress, smallest_first=prefs['max_file_mb'] == 0)
        def analyse(path, notify):
            return analyse_file(path, prefs, self.embedder, semantic, existing, embedded, cancel, notify)
        for path, data in run_files(paths, analyse, workers, cancel, submitted, report_path):
            if cancel.is_set():
                state['cancelled'] = True
                progress(dict(state))
                return state
            key, name = str(path), path.name
            state['path'] = key
            state['errors'] += data['errors']
            if data.get('unchanged'):
                state['unchanged'] += 1
            elif data.get('record'):
                stat, caption_stamp, fingerprint, content, error, segments, parts, vectors = data['record']
                try:
                    # Revalidate at commit, not only when a parallel task finishes.
                    after = path.stat()
                    if (after.st_mtime_ns, after.st_size, after.st_ctime_ns) != (stat.st_mtime_ns, stat.st_size, stat.st_ctime_ns) or caption_fingerprint(path) != caption_stamp:
                        state['errors'] += 1
                    else:
                        with self.connect() as db:
                            db.execute("""INSERT INTO files(path,root,name,size,modified,created,fingerprint,content,error,indexed_at,analyzed_at)
                                VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(path) DO UPDATE SET root=excluded.root,
                                name=excluded.name,size=excluded.size,modified=excluded.modified,created=excluded.created,
                                fingerprint=excluded.fingerprint,content=excluded.content,error=excluded.error,analyzed_at=COALESCE(excluded.analyzed_at,files.analyzed_at)""",
                                (key, str(root), name, stat.st_size, stat.st_mtime, birth_time(path, stat), fingerprint, content, error, time.time(), time.time() if not error and (semantic or content) else None))
                            fid = db.execute("SELECT id FROM files WHERE path=?", (key,)).fetchone()[0]
                            db.execute("DELETE FROM vectors WHERE file_id=?", (fid,))
                            db.execute('DELETE FROM media_segments WHERE file_id=?', (fid,))
                            db.executemany('INSERT INTO media_segments VALUES(?,?,?,?,?)', [(fid, segment.get('part', i), segment['start'], segment['end'], segment['text']) for i, segment in enumerate(segments)])
                            if vectors is not None:
                                db.executemany("INSERT INTO vectors VALUES(?,?,?,?)", [(fid, i, part, vector.astype("<f4").tobytes()) for i, (part, vector) in enumerate(zip(parts, vectors))])
                            self._bump(db)
                        state["updated"] += 1
                except OSError:
                    state['errors'] += 1
            state['processed'] += 1
            progress(dict(state))
        state['cancelled'] = cancel.is_set()
        state["errors"] += len(walk_errors)
        # Never prune after an incomplete traversal (e.g. a disconnected drive).
        try:
            source_online = root.is_dir() and root.stat().st_dev == current_device and volume_identity(root) == current_volume
        except OSError:
            source_online = False
        if not walk_errors and not cancel.is_set() and source_online:
            missing = set(existing) - seen
            with self.connect() as db:
                db.executemany("DELETE FROM files WHERE path=?", [(p,) for p in missing])
                if missing:
                    self._bump(db)
            state["removed"] = len(missing)
        state["cancelled"] = cancel.is_set()
        progress(dict(state))
        return state

    def _paths(self, root, errors, cancel, progress, smallest_first=False):
        own_dir = self.directory.resolve()
        def walk():
            for parent, dirs, names in os.walk(root, followlinks=False, onerror=errors.append):
                dirs[:] = sorted(d for d in dirs if d not in EXCLUDED and not (Path(parent) / d).is_symlink()
                                 and (Path(parent) / d).resolve() != own_dir and Path(parent) / d not in VIRTUAL_ROOTS
                                 and not getattr(Path(parent) / d, "is_junction", lambda: False)())
                for name in sorted(names):
                    if cancel.is_set():
                        return
                    path = Path(parent) / name
                    if path.is_symlink() or not path.is_file() or path.parent.resolve() == own_dir:
                        continue
                    yield path
        # A disk-backed queue avoids retaining millions of paths in worker RAM.
        with tempfile.TemporaryDirectory(prefix="scan-", dir=os.environ.get("FS_WORKER_TMP", self.directory)) as folder:
            queue = sqlite3.connect(Path(folder) / "queue.sqlite3")
            duty = self.preferences()["scan_duty_percent"] / 100
            checkpoint = time.monotonic()
            def throttle():
                nonlocal checkpoint
                if duty < 1:
                    cancel.wait((time.monotonic() - checkpoint) * (1 / duty - 1))
                checkpoint = time.monotonic()
                return int(cancel.is_set())
            queue.set_progress_handler(throttle, 5000)
            try:
                queue.execute("PRAGMA cache_size=-2048")
                queue.execute("PRAGMA temp_store=FILE")
                queue.execute("CREATE TABLE queue(path TEXT PRIMARY KEY,size INTEGER)")
                discovered = 0
                for path in walk():
                    try:
                        queue.execute("INSERT INTO queue VALUES(?,?)", (str(path), path.stat().st_size))
                        discovered += 1
                        if discovered % 100 == 0:
                            queue.commit()
                            progress({"state": "running", "phase": "discovering", "discovered": discovered, "path": str(path), "scanned": 0, "updated": 0})
                    except OSError as exc:
                        errors.append(exc)
                queue.commit()
                if smallest_first:
                    queue.execute("CREATE INDEX queue_size ON queue(size,path)")
                queue.set_progress_handler(None, 0)
                if not cancel.is_set():
                    progress({'phase': 'indexing', 'total': discovered, 'discovered': discovered, 'scanned': 0, 'updated': 0, 'path': str(root)})
                    order = 'size,path' if smallest_first else 'path'
                    for row in queue.execute('SELECT path FROM queue ORDER BY ' + order):
                        yield Path(row[0])
            except sqlite3.OperationalError:
                if not cancel.is_set():
                    raise
            finally:
                queue.close()

    def _vectors(self, db):
        revision = self._meta(db, "revision")
        with self.cache_lock:
            if self.vector_cache is None or self.vector_cache[0] != revision:
                rows = db.execute("SELECT file_id,text,embedding FROM vectors ORDER BY file_id,part").fetchall()
                matrix = np.stack([np.frombuffer(r[2], dtype="<f4") for r in rows]) if rows else np.empty((0, 0), dtype=np.float32)
                self.vector_cache = (revision, rows, matrix)
            return self.vector_cache[1:]

    def search(self, query="", mode="hybrid", limit=50, filename="", root="", modified_after=None, modified_before=None, created_after=None, created_before=None, file_type="", folder="", favourites_only=False, sort_by='score', sort_desc=True, minimum_relevance=0):
        start = time.monotonic()
        if mode not in {"hybrid", "semantic", "text", "filename", "regex"}:
            raise ValueError("Unknown search mode.")
        if len(query) > 1000:
            raise ValueError("Query is too long (maximum 1000 characters).")
        if sort_by not in {'score', 'name', 'path', 'size', 'modified', 'created', 'indexed_at', 'analyzed_at'} or not isinstance(sort_desc, bool):
            raise ValueError('Invalid sort order.')
        minimum_relevance = float(minimum_relevance)
        if not 0 <= minimum_relevance <= 1:
            raise ValueError('Minimum similarity must be between 0 and 1.')
        limit = max(1, min(int(limit), 200))
        clauses, params = [], []
        for column, lower, upper in [("modified", modified_after, modified_before), ("created", created_after, created_before)]:
            for value, op, is_upper in [(lower, ">=", False), (upper, "<", True)]:
                if value:
                    clauses.append(f"{column} {op} ?")
                    params.append(date_bound(value, is_upper))
        if root:
            clauses.append("root=?")
            params.append(str(Path(root).expanduser().resolve()))
        if folder:
            prefix = str(Path(folder).expanduser().resolve()) + os.sep
            clauses.append('substr(path,1,?)=?')
            params.extend([len(prefix), prefix])
        if favourites_only:
            clauses.append("path IN (SELECT path FROM favourites WHERE kind='file')")
        from .media import kind
        if file_type not in {'', 'image', 'video', 'audio', 'document', 'code', 'other'}:
            raise ValueError('Unknown file type.')
        with self.connect() as db:
            # Keep metadata and vectors in one consistent WAL snapshot while indexing.
            db.execute("BEGIN")
            columns = "id,path,root,name,size,modified,created,error,indexed_at,analyzed_at" + (",content" if mode == "regex" else "")
            rows = db.execute("SELECT " + columns + " FROM files" + (" WHERE " + " AND ".join(clauses) if clauses else ""), params).fetchall()
            files = {r["id"]: dict(r) for r in rows if (not filename or filename.casefold() in r["name"].casefold()) and (not file_type or kind(r['name']) == file_type)}
            scores, excerpts = {}, {}
            semantic_error = None
            semantic_scores, literal_ids = {}, set()
            if not query.strip():
                scores = {fid: f["modified"] for fid, f in files.items()}
            elif mode == "filename":
                scores = {fid: 1 for fid, f in files.items() if query.casefold() in f["name"].casefold()}
                literal_ids.update(scores)
            elif mode == "regex":
                try:
                    pattern = regex.compile(query, regex.IGNORECASE | regex.MULTILINE)
                    deadline = time.monotonic() + 2
                    for fid, file in files.items():
                        if time.monotonic() > deadline:
                            raise ValueError("Regex exceeded the 2-second search budget. Narrow your filters.")
                        match = pattern.search(file["name"] + "\n" + file["content"], timeout=0.015)
                        if match:
                            scores[fid] = 1
                            literal_ids.add(fid)
                            excerpts[fid] = (file["name"] + "\n" + file["content"])[max(0, match.start() - 80):match.end() + 180]
                except (regex.error, TimeoutError) as exc:
                    raise ValueError("Invalid or too expensive regular expression.") from exc
            else:
                ranks = []
                if mode in {"text", "hybrid"}:
                    words = regex.findall(r"[\p{L}\p{N}_]+", query)
                    fts_query = " AND ".join('"' + w + '"' for w in words)
                    hits = db.execute("SELECT rowid,bm25(search,8,2,1) AS rank FROM search WHERE search MATCH ? ORDER BY rank", (fts_query,)).fetchall() if words else []
                    text_ids = [r[0] for r in hits if r[0] in files]
                    name_ids = [fid for fid, f in files.items() if query.casefold() in f["name"].casefold()]
                    ranks.append(list(dict.fromkeys(name_ids + text_ids)))
                    literal_ids.update(name_ids + text_ids)
                if mode in {"semantic", "hybrid"}:
                    try:
                        if self._meta(db, "model") not in {None, getattr(self.embedder, "identity", self.embedder.model)}:
                            raise RuntimeError("Index model differs from the selected model. Reindex with the selected model first.")
                        vector_rows, matrix = self._vectors(db)
                        if len(vector_rows):
                            q = self.embedder.embed([query], query=True)[0]
                            if matrix.shape[1] != len(q):
                                raise RuntimeError("Embedding dimensions changed. Rebuild the index.")
                            similarities = matrix @ q
                            for i in np.argsort(-similarities):
                                fid = vector_rows[i][0]
                                if fid in files and fid not in semantic_scores:
                                    semantic_scores[fid] = float(similarities[i])
                                    excerpts[fid] = vector_rows[i][1][:280]
                            ranks.append(list(semantic_scores))
                            if mode == "semantic":
                                scores = semantic_scores
                    except RuntimeError as exc:
                        if mode == "semantic":
                            raise
                        semantic_error = str(exc)
                if mode != "semantic":
                    for ranking in ranks:
                        for rank, fid in enumerate(ranking):
                            scores[fid] = scores.get(fid, 0) + 1 / (60 + rank)
            # RRF is an ordering signal, not an accuracy percentage. Never normalize
            # an unrelated best candidate into a 100% semantic match.
            quality = {fid: 1.0 if fid in literal_ids else max(0.0, min(1.0, semantic_scores.get(fid, 0.0))) for fid in scores}
            if query.strip() and mode in {'hybrid', 'semantic'}:
                scores = {fid: value for fid, value in scores.items() if quality[fid] >= minimum_relevance}
            ranked = sorted(scores, key=lambda fid: files[fid]['path'])
            ranked.sort(key=lambda fid: scores[fid], reverse=True)
            key = 'modified' if sort_by == 'score' and not query.strip() else sort_by
            def sort_value(fid):
                if key == 'score':
                    return quality[fid] if query.strip() else scores[fid]
                value = files[fid][key]
                return value.casefold() if key in {'name', 'path'} else value if value is not None else 0
            ranked.sort(key=sort_value, reverse=sort_desc)
            if key in {'created', 'indexed_at', 'analyzed_at'}:
                ranked.sort(key=lambda fid: files[fid][key] is None)
            devices = {row['root']: (row['device'], row['volume']) for row in db.execute('SELECT root,device,volume FROM locations')}
            results = []
            for fid in ranked[:limit]:
                file = files[fid]
                file.pop("content", None)
                file['availability'] = self.sources.status(file['root'], *devices.get(file['root'], (None, None)))
                file["score"] = round(scores[fid], 5)
                file['relevance'] = None if not query.strip() else round(quality[fid] * 100, 1)
                file['semantic_similarity'] = round(semantic_scores[fid], 5) if fid in semantic_scores else None
                file['relevance_basis'] = 'literal' if fid in literal_ids else 'semantic'
                file['relevance_type'] = 'Exact filename, keyword or regex match; not confidence' if fid in literal_ids else 'Actual semantic cosine similarity; not confidence'
                file["excerpt"] = excerpts.get(fid)
                if file["excerpt"] is None:
                    file["excerpt"] = db.execute("SELECT substr(content,1,280) FROM files WHERE id=?", (fid,)).fetchone()[0]
                results.append(file)
        return {"results": results, "total": len(ranked), "elapsed_ms": round((time.monotonic() - start) * 1000, 1), "warning": semantic_error}

    def remove_root(self, root):
        if not self.index_lock.acquire(blocking=False):
            raise RuntimeError("Wait for indexing to finish before removing a location.")
        try:
            with writer_lock(self.directory / "writer.lock"), self.connect() as db:
                resolved = str(Path(root).expanduser().resolve())
                db.execute("DELETE FROM files WHERE root=?", (resolved,))
                db.execute("DELETE FROM locations WHERE root=?", (resolved,))
                self._bump(db)
        finally:
            self.index_lock.release()
